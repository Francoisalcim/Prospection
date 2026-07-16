#!/usr/bin/env python3
"""
Clinical Trials Prospector - memory-safe version

Key principles:
- Comma-separated user input is split into independent search items.
- Each item may be a simple term or an uppercase Boolean sub-query using AND/OR.
- Lowercase and/or are treated as normal words.
- Searches ClinicalTrials.gov page by page.
- Does not keep raw study JSON in memory after extraction.
- Deduplicates by NCT ID while streaming.
- Applies status at API level; phase/date/country/organization filters locally.
"""

from __future__ import annotations

import csv
import re
import time
from datetime import datetime
from typing import Callable, Dict, List, Optional, Set, Tuple, Any

import requests


ProgressCallback = Optional[Callable[[Dict[str, Any]], None]]


class ClinicalTrialsProspector:
    BASE_URL = "https://clinicaltrials.gov/api/v2/studies"

    ORGANIZATION_TYPES = {
        "university": {
            "keywords": [
                "university", "universite", "universität", "universidad", "universiti",
                "college", "école"
            ],
            "label": "Universities & Colleges",
        },
        "institute": {
            "keywords": ["institute", "institut", "instituto", "research center", "research centre"],
            "label": "Research Institutes",
        },
        "hospital": {
            "keywords": [
                "hospital", "medical center", "medical centre", "health system",
                "health center", "clinic", "clinique", "klinik"
            ],
            "label": "Hospitals & Medical Centers",
        },
        "government": {
            "keywords": [
                "national institutes", "nih", "ministry of health", "department of health",
                "veterans affairs", "va medical", "public health", "government"
            ],
            "label": "Government Agencies",
        },
        "foundation": {
            "keywords": ["foundation", "fondation", "fundacion", "stichting", "trust fund"],
            "label": "Foundations & Trusts",
        },
        "academic": {
            "keywords": ["school of medicine", "school of pharmacy", "faculty of", "academy", "academie", "academic"],
            "label": "Academic Medical Centers",
        },
        "nonprofit": {
            "keywords": ["nonprofit", "non-profit", "charity", "charitable", "society", "association", "organization"],
            "label": "Non-Profit Organizations",
        },
        "company": {
            "keywords": [],
            "label": "Commercial Companies",
        },
    }

    DATA_EXTRACTION_OPTIONS = {
        "sponsors": {"label": "Sponsors & Collaborators", "description": "Lead sponsors and collaborating organizations", "default": True},
        "investigators": {"label": "Principal Investigators", "description": "Lead researchers and their affiliations", "default": False},
        "locations": {"label": "Study Locations", "description": "Facilities, cities, countries where study is conducted", "default": False},
        "interventions": {"label": "Interventions", "description": "Drugs, devices, procedures being tested", "default": False},
        "conditions": {"label": "Conditions", "description": "Diseases and conditions being studied", "default": False},
        "outcomes": {"label": "Study Outcomes", "description": "Primary and secondary outcome measures", "default": False},
        "design": {"label": "Study Design", "description": "Phase, type, enrollment, randomization details", "default": False},
        "eligibility": {"label": "Eligibility Criteria", "description": "Age, gender, inclusion/exclusion criteria", "default": False},
        "contacts": {"label": "Contact Information", "description": "Recruitment contacts and emails", "default": False},
        "timeline": {"label": "Dates & Timeline", "description": "Start date, completion date, last update", "default": False},
    }

    VALID_STATUSES = {
        "RECRUITING", "ACTIVE_NOT_RECRUITING", "NOT_YET_RECRUITING", "COMPLETED",
        "TERMINATED", "WITHDRAWN", "SUSPENDED", "ENROLLING_BY_INVITATION",
        "UNKNOWN"
    }

    VALID_PHASES = {"EARLY_PHASE1", "PHASE1", "PHASE2", "PHASE3", "PHASE4", "NA"}

    VALID_DATE_FIELDS = {"start_date", "completion_date", "last_update"}

    def __init__(
        self,
        include_types: Optional[List[str]] = None,
        exclude_types: Optional[List[str]] = None,
        extraction_options: Optional[List[str]] = None,
    ):
        self.include_types = include_types
        self.exclude_types = exclude_types or []
        self.extraction_options = extraction_options or ["sponsors"]

        # Memory-safe storage: we store only final extracted rows, not raw studies.
        self.extracted_data: List[Dict[str, Any]] = []
        self.seen_nct_ids: Set[str] = set()
        self.processing_errors: List[str] = []
        self.query_terms: List[str] = []
        self.api_requests_made = 0
        self.raw_pages_seen = 0
        self.raw_studies_seen = 0
        self.local_filter_rejections = 0
        self.organization_filter_rejections = 0

    # ---------------------------------------------------------------------
    # Parsing
    # ---------------------------------------------------------------------
    def _split_top_level_commas(self, text: str) -> List[str]:
        """Split on commas, but not commas inside parentheses."""
        parts: List[str] = []
        current: List[str] = []
        depth = 0

        for char in text:
            if char == "(":
                depth += 1
                current.append(char)
            elif char == ")":
                depth = max(0, depth - 1)
                current.append(char)
            elif char == "," and depth == 0:
                part = "".join(current).strip()
                if part:
                    parts.append(part)
                current = []
            else:
                current.append(char)

        part = "".join(current).strip()
        if part:
            parts.append(part)
        return parts

    def clean_query_term(self, term: str) -> str:
        """
        Clean one simple medical term for ClinicalTrials.gov.
        Boolean terms are not passed through this function.
        """
        term = (term or "").strip()
        term = term.replace(" / ", " ")
        term = term.replace("/", " ")
        term = term.replace("&", " and ")
        term = term.replace("(", " ").replace(")", " ")
        term = re.sub(r"\s+", " ", term).strip()
        return term

    def parse_keyword_terms(self, keyword_string: str) -> List[str]:
        """
        Convert user input into one or more API query terms.

        Rules:
        1. Split first on top-level commas.
        2. Each item may be:
           - a simple medical term
           - a Boolean sub-query using uppercase AND/OR
        3. Lowercase and/or remain normal text.
        """
        expression = (keyword_string or "").strip()
        if not expression:
            return []

        expression = re.sub(r"\s+", " ", expression)
        raw_items = self._split_top_level_commas(expression)
        query_terms: List[str] = []

        for item in raw_items:
            item = item.strip()
            if not item:
                continue

            has_boolean = re.search(r"\b(?:AND|OR)\b", item) is not None
            if has_boolean:
                query = self.parse_keyword_expression(item)
            else:
                query = self.clean_query_term(item)

            if query:
                query_terms.append(query)

        return query_terms

    def parse_keyword_expression(self, keyword_string: str) -> str:
        """Validate one Boolean sub-query. Should not receive comma-separated full input."""
        expression = (keyword_string or "").strip()
        if not expression:
            return ""
        expression = re.sub(r"\s+", " ", expression)

        if "," in expression:
            raise ValueError(
                "Internal parser error: Boolean expression received commas. "
                "The full input must go through parse_keyword_terms() first."
            )

        self._validate_keyword_expression(expression)
        return expression

    def _validate_keyword_expression(self, expression: str) -> None:
        balance = 0
        for char in expression:
            if char == "(":
                balance += 1
            elif char == ")":
                balance -= 1
            if balance < 0:
                raise ValueError("Invalid Boolean expression: closing parenthesis before opening parenthesis.")
        if balance != 0:
            raise ValueError("Invalid Boolean expression: unbalanced parentheses.")

        if re.search(r"\b(?:AND|OR)\s+(?:AND|OR)\b", expression):
            raise ValueError("Invalid Boolean expression: two operators are next to each other.")
        if re.search(r"^\s*(?:AND|OR)\b", expression) or re.search(r"\b(?:AND|OR)\s*$", expression):
            raise ValueError("Invalid Boolean expression: expression cannot start or end with AND/OR.")

    # ---------------------------------------------------------------------
    # Validation
    # ---------------------------------------------------------------------
    def validate_filters(self, statuses=None, phases=None, date_field=None) -> None:
        if statuses:
            invalid = set(statuses) - self.VALID_STATUSES
            if invalid:
                raise ValueError(f"Invalid status value(s): {', '.join(sorted(invalid))}")
        if phases:
            invalid = set(phases) - self.VALID_PHASES
            if invalid:
                raise ValueError(f"Invalid phase value(s): {', '.join(sorted(invalid))}")
        if date_field and date_field not in self.VALID_DATE_FIELDS:
            raise ValueError(f"Invalid date field: {date_field}")

    # ---------------------------------------------------------------------
    # Organization classification
    # ---------------------------------------------------------------------
    def get_organization_type(self, name: str) -> str:
        if not name:
            return "unknown"
        name_lower = name.lower()
        for org_type, info in self.ORGANIZATION_TYPES.items():
            if org_type == "company":
                continue
            for keyword in info["keywords"]:
                if keyword in name_lower:
                    return org_type
        return "company"

    def should_include_organization(self, name: str) -> bool:
        if not name:
            return False
        org_type = self.get_organization_type(name)
        if self.include_types is not None:
            return org_type in self.include_types
        return org_type not in self.exclude_types

    # ---------------------------------------------------------------------
    # Local filters
    # ---------------------------------------------------------------------
    def _extract_study_phases(self, study: Dict[str, Any]) -> List[str]:
        return (
            study.get("protocolSection", {})
            .get("designModule", {})
            .get("phases", [])
        ) or []

    def _extract_study_countries(self, study: Dict[str, Any]) -> Set[str]:
        locations = (
            study.get("protocolSection", {})
            .get("contactsLocationsModule", {})
            .get("locations", [])
        ) or []
        return {
            str(loc.get("country", "")).strip().lower()
            for loc in locations
            if loc.get("country")
        }

    def _extract_study_date(self, study: Dict[str, Any], date_field: str) -> Optional[str]:
        status_module = study.get("protocolSection", {}).get("statusModule", {})
        field_map = {
            "start_date": "startDateStruct",
            "completion_date": "completionDateStruct",
            "last_update": "lastUpdatePostDateStruct",
        }
        struct_name = field_map.get(date_field)
        if not struct_name:
            return None
        return status_module.get(struct_name, {}).get("date")

    def _normalize_date_for_compare(self, value: str, end_of_period: bool = False) -> Optional[str]:
        """
        ClinicalTrials.gov dates can be YYYY, YYYY-MM, or YYYY-MM-DD.
        Normalize for lexical YYYY-MM-DD comparison.
        """
        if not value:
            return None
        value = value.strip()
        if re.fullmatch(r"\d{4}", value):
            return f"{value}-12-31" if end_of_period else f"{value}-01-01"
        if re.fullmatch(r"\d{4}-\d{2}", value):
            return f"{value}-31" if end_of_period else f"{value}-01"
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            return value
        return value[:10]

    def _date_in_range(self, date_value: Optional[str], date_from: Optional[str], date_to: Optional[str]) -> bool:
        if not date_value:
            return False
        lower = self._normalize_date_for_compare(date_value, end_of_period=False)
        upper = self._normalize_date_for_compare(date_value, end_of_period=True)
        if not lower or not upper:
            return False
        if date_from and upper < date_from:
            return False
        if date_to and lower > date_to:
            return False
        return True

    def _passes_local_filters(
        self,
        study: Dict[str, Any],
        phases: Optional[List[str]] = None,
        date_field: Optional[str] = None,
        date_from: Optional[str] = None,
        date_to: Optional[str] = None,
        countries: Optional[List[str]] = None,
    ) -> bool:
        if phases:
            study_phases = set(self._extract_study_phases(study))
            if not set(phases).intersection(study_phases):
                return False

        if date_field and (date_from or date_to):
            study_date = self._extract_study_date(study, date_field)
            if not self._date_in_range(study_date, date_from, date_to):
                return False

        if countries:
            requested = {c.strip().lower() for c in countries if c.strip()}
            study_countries = self._extract_study_countries(study)
            if not requested.intersection(study_countries):
                return False

        return True

    # ---------------------------------------------------------------------
    # Fetching, streaming extraction, deduplication
    # ---------------------------------------------------------------------
    def fetch_trials(
        self,
        keywords: str,
        statuses: Optional[List[str]] = None,
        phases: Optional[List[str]] = None,
        max_results: Optional[int] = 500,
        progress_callback: ProgressCallback = None,
        date_field: Optional[str] = None,
        date_from: Optional[str] = None,
        date_to: Optional[str] = None,
        countries: Optional[List[str]] = None,
        request_delay_seconds: float = 0.15,
        page_size: int = 100,
    ) -> List[Dict[str, Any]]:
        """
        Fetch and extract studies page by page.
        Returns extracted rows only. Does not store raw study JSON.
        """
        self.extracted_data = []
        self.seen_nct_ids = set()
        self.processing_errors = []
        self.api_requests_made = 0
        self.raw_pages_seen = 0
        self.raw_studies_seen = 0
        self.local_filter_rejections = 0
        self.organization_filter_rejections = 0

        self.validate_filters(statuses=statuses, phases=phases, date_field=date_field)
        self.query_terms = self.parse_keyword_terms(keywords)
        if not self.query_terms:
            return []

        target_limit = max_results

        if progress_callback:
            progress_callback({
                "stage": "starting",
                "query_count": len(self.query_terms),
                "max_results": target_limit,
                "extracted": 0,
            })

        for query_index, query_term in enumerate(self.query_terms, start=1):
            page_token = None

            while target_limit is None or len(self.extracted_data) < target_limit:
                params = {
                    "query.term": query_term,
                    "pageSize": page_size,
                    "format": "json",
                    "countTotal": "true",
                }
                if statuses:
                    params["filter.overallStatus"] = ",".join(statuses)
                # Phase/date/countries are local filters. Do not send filter.phase.
                if page_token:
                    params["pageToken"] = page_token

                response = requests.get(self.BASE_URL, params=params, timeout=90)
                self.api_requests_made += 1

                if response.status_code != 200:
                    message = (
                        f"ClinicalTrials.gov API error {response.status_code} for query '{query_term}'. "
                        f"URL: {response.url}. Response: {response.text[:500]}"
                    )
                    self.processing_errors.append(message)
                    response.raise_for_status()

                data = response.json()
                self.raw_pages_seen += 1
                studies = data.get("studies", []) or []
                self.raw_studies_seen += len(studies)

                if not studies:
                    break

                for study in studies:
                    nct_id = (
                        study.get("protocolSection", {})
                        .get("identificationModule", {})
                        .get("nctId", "")
                    )
                    if not nct_id or nct_id in self.seen_nct_ids:
                        continue

                    if not self._passes_local_filters(
                        study,
                        phases=phases,
                        date_field=date_field,
                        date_from=date_from,
                        date_to=date_to,
                        countries=countries,
                    ):
                        self.local_filter_rejections += 1
                        continue

                    row = self.extract_one_study(study)
                    if not row:
                        continue

                    self.extracted_data.append(row)
                    self.seen_nct_ids.add(nct_id)

                    if target_limit is not None and len(self.extracted_data) >= target_limit:
                        break

                if progress_callback:
                    progress_callback({
                        "stage": "fetching",
                        "current_query": query_term,
                        "query_index": query_index,
                        "query_count": len(self.query_terms),
                        "raw_pages_seen": self.raw_pages_seen,
                        "raw_studies_seen": self.raw_studies_seen,
                        "extracted": len(self.extracted_data),
                        "deduplicated_nct_ids": len(self.seen_nct_ids),
                        "local_filter_rejections": self.local_filter_rejections,
                        "organization_filter_rejections": self.organization_filter_rejections,
                        "api_requests_made": self.api_requests_made,
                        "has_next_page": bool(data.get("nextPageToken")),
                        "max_results": target_limit,
                    })

                page_token = data.get("nextPageToken")
                if not page_token:
                    break

                time.sleep(request_delay_seconds)

            if target_limit is not None and len(self.extracted_data) >= target_limit:
                break

        if progress_callback:
            progress_callback({
                "stage": "completed",
                "extracted": len(self.extracted_data),
                "deduplicated_nct_ids": len(self.seen_nct_ids),
                "raw_studies_seen": self.raw_studies_seen,
                "api_requests_made": self.api_requests_made,
            })

        return self.extracted_data

    def extract_data(self, trials: Optional[List[Dict[str, Any]]] = None) -> List[Dict[str, Any]]:
        """
        Backward-compatible method.
        In memory-safe mode, fetch_trials already extracts data.
        If trials are provided, extract them one by one.
        """
        if trials is None:
            return self.extracted_data
        self.extracted_data = []
        self.seen_nct_ids = set()
        for study in trials:
            row = self.extract_one_study(study)
            if row:
                self.extracted_data.append(row)
        return self.extracted_data

    # ---------------------------------------------------------------------
    # Extraction
    # ---------------------------------------------------------------------
    def extract_one_study(self, study: Dict[str, Any]) -> Optional[Dict[str, Any]]:
        try:
            ps = study.get("protocolSection", {})
            nct_id = ps.get("identificationModule", {}).get("nctId", "")
            if not nct_id:
                return None

            extracted: Dict[str, Any] = {
                "nct_id": nct_id,
                "title": ps.get("identificationModule", {}).get("briefTitle", ""),
                "status": ps.get("statusModule", {}).get("overallStatus", ""),
            }

            if "sponsors" in self.extraction_options:
                extracted.update(self._extract_sponsors(ps))
            if "investigators" in self.extraction_options:
                extracted.update(self._extract_investigators(ps))
            if "locations" in self.extraction_options:
                extracted.update(self._extract_locations(ps))
            if "interventions" in self.extraction_options:
                extracted.update(self._extract_interventions(ps))
            if "conditions" in self.extraction_options:
                extracted.update(self._extract_conditions(ps))
            if "outcomes" in self.extraction_options:
                extracted.update(self._extract_outcomes(ps))
            if "design" in self.extraction_options:
                extracted.update(self._extract_design(ps))
            if "eligibility" in self.extraction_options:
                extracted.update(self._extract_eligibility(ps))
            if "contacts" in self.extraction_options:
                extracted.update(self._extract_contacts(ps))
            if "timeline" in self.extraction_options:
                extracted.update(self._extract_timeline(ps))

            if "sponsors" in self.extraction_options:
                lead_sponsor = extracted.get("lead_sponsor", "")
                if lead_sponsor and not self.should_include_organization(lead_sponsor):
                    self.organization_filter_rejections += 1
                    return None

            return extracted
        except Exception as exc:
            self.processing_errors.append(f"Error extracting study: {exc}")
            return None

    def _extract_sponsors(self, ps: Dict[str, Any]) -> Dict[str, Any]:
        sponsors_module = ps.get("sponsorsCollaboratorsModule") or ps.get("sponsorCollaboratorsModule", {})
        lead_sponsor = sponsors_module.get("leadSponsor", {}).get("name", "")
        lead_sponsor_class = sponsors_module.get("leadSponsor", {}).get("class", "")
        collaborators = sponsors_module.get("collaborators", []) or []
        collab_names = [c.get("name", "") for c in collaborators if c.get("name")]
        return {
            "lead_sponsor": lead_sponsor,
            "lead_sponsor_class": lead_sponsor_class,
            "lead_sponsor_type": self.get_organization_type(lead_sponsor),
            "collaborators": "; ".join(collab_names),
            "collaborator_count": len(collab_names),
        }

    def _extract_investigators(self, ps: Dict[str, Any]) -> Dict[str, Any]:
        contacts_module = ps.get("contactsLocationsModule", {})
        officials = contacts_module.get("overallOfficials", []) or []
        pi_names, pi_affiliations = [], []
        for official in officials:
            if official.get("role") in ["PRINCIPAL_INVESTIGATOR", "STUDY_DIRECTOR"]:
                if official.get("name"):
                    pi_names.append(official.get("name", ""))
                if official.get("affiliation"):
                    pi_affiliations.append(official.get("affiliation", ""))
        return {
            "principal_investigators": "; ".join(pi_names),
            "pi_affiliations": "; ".join(pi_affiliations),
            "pi_count": len(pi_names),
        }

    def _extract_locations(self, ps: Dict[str, Any]) -> Dict[str, Any]:
        contacts_module = ps.get("contactsLocationsModule", {})
        locations = contacts_module.get("locations", []) or []
        facilities, cities, countries = set(), set(), set()
        for loc in locations:
            if loc.get("facility"):
                facilities.add(loc["facility"])
            if loc.get("city"):
                cities.add(loc["city"])
            if loc.get("country"):
                countries.add(loc["country"])
        return {
            "facilities": "; ".join(sorted(facilities)),
            "cities": "; ".join(sorted(cities)),
            "countries": "; ".join(sorted(countries)),
            "location_count": len(locations),
        }

    def _extract_interventions(self, ps: Dict[str, Any]) -> Dict[str, Any]:
        arms_module = ps.get("armsInterventionsModule", {})
        interventions = arms_module.get("interventions", []) or []
        drugs, devices, procedures, other = [], [], [], []
        for intervention in interventions:
            int_type = intervention.get("type", "")
            name = intervention.get("name", "")
            if not name:
                continue
            if int_type == "DRUG":
                drugs.append(name)
            elif int_type == "DEVICE":
                devices.append(name)
            elif int_type in ["PROCEDURE", "SURGERY"]:
                procedures.append(name)
            else:
                other.append(name)
        return {
            "drugs": "; ".join(drugs),
            "devices": "; ".join(devices),
            "procedures": "; ".join(procedures),
            "other_interventions": "; ".join(other),
            "intervention_count": len(interventions),
        }

    def _extract_conditions(self, ps: Dict[str, Any]) -> Dict[str, Any]:
        cond_module = ps.get("conditionsModule", {})
        conditions = cond_module.get("conditions", []) or []
        keywords = cond_module.get("keywords", []) or []
        return {
            "conditions": "; ".join(conditions),
            "keywords": "; ".join(keywords),
            "condition_count": len(conditions),
        }

    def _extract_outcomes(self, ps: Dict[str, Any]) -> Dict[str, Any]:
        outcomes_module = ps.get("outcomesModule", {})
        primary_outcomes = outcomes_module.get("primaryOutcomes", []) or []
        secondary_outcomes = outcomes_module.get("secondaryOutcomes", []) or []
        primary = [o.get("measure", "") for o in primary_outcomes if o.get("measure")]
        secondary = [o.get("measure", "") for o in secondary_outcomes if o.get("measure")]
        return {
            "primary_outcomes": "; ".join(primary),
            "secondary_outcomes": "; ".join(secondary),
            "primary_outcome_count": len(primary),
            "secondary_outcome_count": len(secondary),
        }

    def _extract_design(self, ps: Dict[str, Any]) -> Dict[str, Any]:
        design_module = ps.get("designModule", {})
        phases = design_module.get("phases", []) or []
        enrollment = design_module.get("enrollmentInfo", {}) or {}
        design_info = design_module.get("designInfo", {}) or {}
        return {
            "phase": ", ".join(phases),
            "study_type": design_module.get("studyType", ""),
            "enrollment": enrollment.get("count", 0),
            "allocation": design_info.get("allocation", ""),
            "intervention_model": design_info.get("interventionModel", ""),
            "primary_purpose": design_info.get("primaryPurpose", ""),
            "masking": (design_info.get("maskingInfo", {}) or {}).get("masking", ""),
        }

    def _extract_eligibility(self, ps: Dict[str, Any]) -> Dict[str, Any]:
        eligibility_module = ps.get("eligibilityModule", {})
        return {
            "min_age": eligibility_module.get("minimumAge", ""),
            "max_age": eligibility_module.get("maximumAge", ""),
            "sex": eligibility_module.get("sex", ""),
            "healthy_volunteers": eligibility_module.get("healthyVolunteers", ""),
            "eligibility_criteria": (eligibility_module.get("eligibilityCriteria", "") or "")[:1000],
        }

    def _extract_contacts(self, ps: Dict[str, Any]) -> Dict[str, Any]:
        contacts_module = ps.get("contactsLocationsModule", {})
        central_contacts = contacts_module.get("centralContacts", []) or []
        names, emails, phones = [], [], []
        for contact in central_contacts:
            if contact.get("name"):
                names.append(contact["name"])
            if contact.get("email"):
                emails.append(contact["email"])
            if contact.get("phone"):
                phones.append(contact["phone"])
        return {
            "contact_name": "; ".join(names),
            "contact_email": "; ".join(emails),
            "contact_phone": "; ".join(phones),
        }

    def _extract_timeline(self, ps: Dict[str, Any]) -> Dict[str, Any]:
        status_module = ps.get("statusModule", {})
        return {
            "start_date": status_module.get("startDateStruct", {}).get("date", ""),
            "completion_date": status_module.get("completionDateStruct", {}).get("date", ""),
            "last_update": status_module.get("lastUpdatePostDateStruct", {}).get("date", ""),
        }

    # ---------------------------------------------------------------------
    # Export
    # ---------------------------------------------------------------------
    def export_to_xlsx(self, filename: Optional[str] = None, column_order: Optional[List[str]] = None) -> Optional[str]:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            print("openpyxl not installed. Run: pip install openpyxl")
            return None

        if not self.extracted_data:
            print("No data to export")
            return None

        if filename is None:
            date = datetime.now().strftime("%Y-%m-%d_%H%M%S")
            filename = f"ClinicalTrials_Export_{date}.xlsx"

        all_keys: Set[str] = set()
        for row in self.extracted_data:
            all_keys.update(row.keys())

        if column_order:
            fieldnames = [col for col in column_order if col in all_keys]
            remaining = [col for col in sorted(all_keys) if col not in fieldnames]
            fieldnames.extend(remaining)
        else:
            fieldnames = sorted(all_keys)
            for field in reversed(["nct_id", "title", "status", "lead_sponsor"]):
                if field in fieldnames:
                    fieldnames.remove(field)
                    fieldnames.insert(0, field)

        wb = Workbook()
        ws = wb.active
        ws.title = "Clinical Trials Data"

        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx, fieldname in enumerate(fieldnames, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = fieldname.replace("_", " ").title()
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        wrap_fields = {
            "title", "eligibility_criteria", "conditions", "interventions", "collaborators",
            "primary_outcomes", "secondary_outcomes", "facilities"
        }
        for row_idx, record in enumerate(self.extracted_data, 2):
            for col_idx, fieldname in enumerate(fieldnames, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = record.get(fieldname, "")
                if fieldname in wrap_fields:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

        for col_idx, fieldname in enumerate(fieldnames, 1):
            col = get_column_letter(col_idx)
            if fieldname == "nct_id":
                width = 14
            elif fieldname == "title":
                width = 55
            elif fieldname in wrap_fields:
                width = 45
            elif fieldname in ["lead_sponsor", "principal_investigators"]:
                width = 32
            else:
                width = 18
            ws.column_dimensions[col].width = width

        ws.freeze_panes = "A2"
        wb.save(filename)
        return filename

    def export_to_csv(self, filename: Optional[str] = None) -> Optional[str]:
        if not self.extracted_data:
            return None
        if filename is None:
            date = datetime.now().strftime("%Y-%m-%d_%H%M%S")
            filename = f"ClinicalTrials_Export_{date}.csv"
        all_keys: Set[str] = set()
        for row in self.extracted_data:
            all_keys.update(row.keys())
        fieldnames = sorted(all_keys)
        for field in reversed(["nct_id", "title", "status", "lead_sponsor"]):
            if field in fieldnames:
                fieldnames.remove(field)
                fieldnames.insert(0, field)
        with open(filename, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(self.extracted_data)
        return filename
